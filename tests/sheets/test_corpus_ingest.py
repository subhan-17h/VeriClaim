"""PostgreSQL acceptance tests for the generated spreadsheet corpus."""

from __future__ import annotations

from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from psycopg import sql

from vericlaim.config import Settings
from vericlaim.corpus.spreadsheets import generate_spreadsheet_corpus
from vericlaim.sheets.ingest import LINEAGE_COLUMNS, ingest_workbook
from vericlaim.sheets.profiler import profile_workbook
from vericlaim.sql.contexts import load_contexts
from vericlaim.sql.db import Database

live = pytest.mark.postgres


def _drop_sheet_tables(database: Database) -> None:
    with database.connection() as conn:
        conn.execute(
            "DO $$ DECLARE t text; BEGIN "
            "FOR t IN SELECT tablename FROM pg_tables WHERE schemaname = 'sheets' "
            "LOOP EXECUTE format('DROP TABLE IF EXISTS sheets.%I CASCADE', t); "
            "END LOOP; END $$;"
        )


@pytest.fixture(scope="module")
def admin(settings: Settings) -> Iterator[Database]:
    database = Database(settings.dsn(readonly=False), statement_timeout_ms=15_000)
    _drop_sheet_tables(database)
    yield database
    _drop_sheet_tables(database)
    database.close()


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, ...]:
    output = tmp_path_factory.mktemp("spreadsheet-corpus")
    return tuple(generate_spreadsheet_corpus(output, seed=42))


@pytest.fixture(scope="module")
def ingested(admin: Database, generated: tuple[Path, ...]):
    return tuple(
        table for path in generated for table in ingest_workbook(admin, path)
    )


@live
def test_the_six_workbooks_create_exactly_the_six_declared_tables(
    admin: Database, ingested, settings: Settings
) -> None:
    declared = set(load_contexts(settings.sheets_context_dir))
    with admin.connection() as conn:
        actual = {
            f"sheets.{row[0]}"
            for row in conn.execute(
                "SELECT tablename FROM pg_tables WHERE schemaname = 'sheets'"
            ).fetchall()
        }

    assert {table.qualified for table in ingested} == declared
    assert actual == declared


@live
def test_every_ingested_row_has_complete_lineage_and_excludes_footers(
    admin: Database, generated: tuple[Path, ...], ingested
) -> None:
    path_by_name = {path.name: path for path in generated}
    for table in ingested:
        profile = profile_workbook(path_by_name[table.workbook])[0].tables[0]
        expected_rows = profile.last_data_row - profile.first_data_row + 1
        lineage_present = sql.SQL(" AND ").join(
            sql.SQL("{} IS NOT NULL").format(sql.Identifier(column))
            for column in LINEAGE_COLUMNS
        )
        with admin.connection() as conn:
            count = conn.execute(
                sql.SQL("SELECT count(*) FROM {}").format(
                    sql.Identifier("sheets", table.table)
                )
            ).fetchone()[0]
            complete = conn.execute(
                sql.SQL("SELECT count(*) FROM {} WHERE {}").format(
                    sql.Identifier("sheets", table.table), lineage_present
                )
            ).fetchone()[0]
            values = conn.execute(
                sql.SQL("SELECT * FROM {}").format(
                    sql.Identifier("sheets", table.table)
                )
            ).fetchall()

        assert table.row_count == expected_rows
        assert count == expected_rows
        assert complete == expected_rows
        assert all(value != "TOTAL" for row in values for value in row)


@live
def test_a_known_percentage_round_trips_to_the_cell_named_by_its_a1_range(
    admin: Database, generated: tuple[Path, ...], ingested
) -> None:
    table = next(
        item
        for item in ingested
        if item.workbook == "Regional_Inspection_Compliance_Q1.xlsx"
    )
    with admin.connection() as conn:
        stored_rate, locator = conn.execute(
            sql.SQL(
                "SELECT compliance_rate, _a1_range FROM {} WHERE region = %s"
            ).format(sql.Identifier("sheets", table.table)),
            ("Lahore Central",),
        ).fetchone()

    path = next(path for path in generated if path.name == table.workbook)
    book = load_workbook(path, data_only=True)
    try:
        sheet = book[table.sheet]
        min_column, min_row, max_column, max_row = range_boundaries(locator)
        assert (min_column, min_row, max_column, max_row) == (1, 4, 5, 4)
        assert sheet["A4"].value == "Lahore Central"
        assert sheet["D4"].value == "60%"
    finally:
        book.close()

    assert stored_rate == Decimal("0.6")
